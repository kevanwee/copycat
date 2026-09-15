import io
import random
from datetime import UTC, datetime, timedelta
from pathlib import Path
import pytest
from pypdf import PdfReader, PdfWriter
from app.services.similarity.text_similarity import compute_text_similarity, lcs_ratio


def new_case(client, intake=None):
    response = client.post('/api/v1/cases', json={'intake': intake or {}})
    assert response.status_code == 200, response.text
    body = response.json()
    client.headers['X-Case-Token'] = body['access_token']
    return body['case_id'], body['access_token']


def upload_pair(client, case, content=b'An original explanation of copyright and creative expression in Singapore.', media='text', extension='txt'):
    for role in ('original', 'alleged'):
        response = client.post(f'/api/v1/cases/{case}/artifacts', data={'role': role, 'media_type': media}, files={'file': (role+'.'+extension, content)})
        assert response.status_code == 200, response.text


def test_access_is_case_scoped(client):
    a, key_a = new_case(client)
    upload_pair(client, a)
    job = client.post(f'/api/v1/cases/{a}/analyze').json()['job_id']
    _, key_b = new_case(client)
    for url in [f'/api/v1/cases/{a}', f'/api/v1/cases/{a}/report', f'/api/v1/cases/{a}/report.pdf', f'/api/v1/jobs/{job}']:
        assert client.get(url).status_code == 404
    assert client.delete(f'/api/v1/cases/{a}').status_code == 404
    client.headers['X-Case-Token'] = key_a
    response = client.get(f'/api/v1/cases/{a}/report')
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert 'access_hash' not in response.text


def test_reproducible_reports_and_revision_invalidation(client):
    ids = []
    for _ in range(2):
        case, _ = new_case(client)
        upload_pair(client, case)
        assert client.post(f'/api/v1/cases/{case}/analyze').status_code == 200
        report = client.get(f'/api/v1/cases/{case}/report').json()['report']
        ids.append(report['report_id'])
    assert ids[0] == ids[1]
    pdf = client.get(f'/api/v1/cases/{case}/report.pdf')
    assert pdf.content.startswith(b'%PDF-')
    text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert 'sso.agc.gov.sg' in text and 'Recorded basis' in text
    response = client.put(f'/api/v1/cases/{case}/intake', json={'title': 'Revised <b>title</b>'})
    assert response.status_code == 200
    assert client.get(f'/api/v1/cases/{case}/report').status_code == 409
    assert client.get(f'/api/v1/cases/{case}/report.pdf').status_code == 409
    client.post(f'/api/v1/cases/{case}/analyze')
    revised = client.get(f'/api/v1/cases/{case}/report').json()['report']
    assert revised['report_id'] != ids[-1]


def test_duplicate_run_and_mutation_rejected(client, monkeypatch):
    from app.api.routes import cases
    from app.db.session import SessionLocal
    from app.db.models import Case, Job
    case, _ = new_case(client)
    upload_pair(client, case)
    monkeypatch.setattr(cases, 'run_case_analysis', lambda **kwargs: None)
    assert client.post(f'/api/v1/cases/{case}/analyze').status_code == 200
    assert client.post(f'/api/v1/cases/{case}/analyze').status_code == 409
    assert client.put(f'/api/v1/cases/{case}/intake', json={}).status_code == 409
    assert client.delete(f'/api/v1/cases/{case}').status_code == 409
    with SessionLocal() as db:
        db.query(Job).filter_by(case_id=case).update({'status': 'failed'})
        db.query(Case).filter_by(id=case).update({'status': 'failed'})
        db.commit()


@pytest.mark.parametrize('filename,content', [('empty.txt', b''), ('bad.pdf', b'not a pdf'), ('bad.docx', b'not a zip'), ('binary.txt', b'\xff\x00'), ('spaces.txt', b'  '), ('script.html', b'<h1>test</h1>')])
def test_invalid_uploads(client, filename, content):
    case, _ = new_case(client)
    result = client.post(f'/api/v1/cases/{case}/artifacts', data={'role': 'original', 'media_type': 'text'}, files={'file': (filename, content)})
    assert result.status_code == 422


def test_oversized_upload(client, monkeypatch):
    from app.api.routes.cases import settings
    monkeypatch.setattr(settings, 'max_text_mb', 1)
    case, _ = new_case(client)
    result = client.post(f'/api/v1/cases/{case}/artifacts', data={'role': 'original', 'media_type': 'text'}, files={'file': ('large.txt', b'a' * (1024*1024+1))})
    assert result.status_code == 413


def test_ingress_limit_and_early_upload_auth(client):
    case, key = new_case(client)
    result = client.post('/api/v1/cases', content=b' ' * (256*1024+1), headers={'Content-Type': 'application/json'})
    assert result.status_code == 413
    result = client.post(f'/api/v1/cases/{case}/artifacts', content=b'not multipart', headers={'X-Case-Token': 'wrong'})
    assert result.status_code == 404


def test_docx_tables_are_extracted(tmp_path):
    from docx import Document
    from app.services.extraction.text import extract_text
    doc = Document(); doc.add_paragraph('Before'); table = doc.add_table(rows=1, cols=1); table.cell(0,0).text = 'Inside the table'; doc.add_paragraph('After')
    path = tmp_path / 'table.docx'; doc.save(path)
    text = extract_text(path).text
    assert text.index('Before') < text.index('Inside the table') < text.index('After')


def test_unreadable_pdf_fails_without_false_score(client):
    writer = PdfWriter(); writer.add_blank_page(100, 100)
    buf = io.BytesIO(); writer.write(buf)
    case, _ = new_case(client)
    upload_pair(client, case, buf.getvalue(), extension='pdf')
    job = client.post(f'/api/v1/cases/{case}/analyze').json()['job_id']
    result = client.get(f'/api/v1/jobs/{job}').json()
    assert result['status'] == 'failed'
    assert 'No readable words' in result['error']
    assert 'Traceback' not in result['error']
    assert client.get(f'/api/v1/cases/{case}/report').status_code == 409


def test_replacement_and_deletion_remove_files(client):
    from app.db.session import SessionLocal
    from app.db.models import Artifact, Case, CaseReport, SimilarityMetric
    from app.core.config import get_settings
    case, _ = new_case(client)
    upload_pair(client, case)
    with SessionLocal() as db:
        paths = [Path(a.storage_path) for a in db.query(Artifact).filter_by(case_id=case)]
    upload_pair(client, case, b'Different comparison content with enough words.')
    assert not any(p.exists() for p in paths)
    client.post(f'/api/v1/cases/{case}/analyze')
    assert client.delete(f'/api/v1/cases/{case}').status_code == 204
    assert not (get_settings().storage_root / case).exists()
    assert not (get_settings().report_root / f'{case}.pdf').exists()
    with SessionLocal() as db:
        assert db.get(Case, case) is None
        assert db.query(CaseReport).filter_by(case_id=case).count() == 0
        assert db.query(SimilarityMetric).filter_by(case_id=case).count() == 0


def test_expiry_covers_report_text_and_derived_files(client):
    from app.db.session import SessionLocal
    from app.db.models import Case
    from app.core.config import get_settings
    from app.services.retention import cleanup_expired
    case, _ = new_case(client)
    upload_pair(client, case)
    client.post(f'/api/v1/cases/{case}/analyze')
    derived = get_settings().report_root / case / 'frames'
    derived.mkdir(parents=True); (derived / 'test.jpg').write_bytes(b'derived')
    with SessionLocal() as db:
        db.get(Case, case).created_at = datetime.now(UTC)-timedelta(hours=25)
        db.commit()
    assert client.get(f'/api/v1/cases/{case}/report').status_code == 410
    with SessionLocal() as db:
        assert cleanup_expired(db) >= 1
        assert db.get(Case, case) is None
    assert not derived.exists()


def test_short_and_empty_text_regressions():
    for a, b in [('', ''), ('!!!', '???'), ('good text', '')]:
        with pytest.raises(ValueError): compute_text_similarity(a, b)
    assert compute_text_similarity('alpha beta', 'carrot potato').headline_score == 0
    assert compute_text_similarity('tiny text', 'tiny text').headline_score == 1
    assert compute_text_similarity('你好世界', '你好世界').headline_score == 1
    assert compute_text_similarity('café story', 'café story').headline_score == 1


def test_repetition_coverage_and_excerpt_containment():
    result = compute_text_similarity('word ' * 10000, 'word ' * 10000)
    assert result.headline_score == 1
    assert result.coverage['original'] == result.coverage['alleged'] == 1
    assert len(result.matched_passages) == 1
    excerpt = 'a distinctive and carefully composed passage of important original words'
    result = compute_text_similarity('filler ' * 5000 + excerpt, excerpt)
    assert result.coverage['alleged'] == 1
    assert result.coverage['original'] < .01


def test_exact_lcs_against_reference():
    rng = random.Random(8)
    for _ in range(50):
        a = [rng.choice('abc') for _ in range(rng.randrange(1, 20))]
        b = [rng.choice('abc') for _ in range(rng.randrange(1, 20))]
        prev = [0] * (len(b)+1)
        for x in a:
            row = [0]
            for j, y in enumerate(b): row.append(prev[j]+1 if x == y else max(prev[j+1], row[-1]))
            prev = row
        assert lcs_ratio(a, b) == prev[-1] / max(len(a), len(b))


def test_image_flow_and_preview(client):
    pytest.importorskip('imagehash'); pytest.importorskip('cv2'); pytest.importorskip('skimage')
    from PIL import Image
    buf = io.BytesIO(); Image.new('RGB', (64, 64), 'green').save(buf, format='PNG')
    case, _ = new_case(client)
    upload_pair(client, case, buf.getvalue(), media='image', extension='png')
    client.post(f'/api/v1/cases/{case}/analyze')
    report = client.get(f'/api/v1/cases/{case}/report').json()['report']
    assert report['headline_overlap_percentage'] == 100
    assert 'paths' not in report['evidence']
    assert client.get(f'/api/v1/cases/{case}/preview/original').headers['content-type'] == 'image/png'


def test_visual_video_flow(client, tmp_path):
    import shutil
    import subprocess
    from app.core.config import get_settings
    settings = get_settings()
    if not shutil.which(settings.ffmpeg_bin) or not shutil.which(settings.ffprobe_bin):
        pytest.skip('FFmpeg/ffprobe not installed')
    pytest.importorskip('imagehash'); pytest.importorskip('cv2'); pytest.importorskip('skimage')
    path = tmp_path / 'synthetic.mp4'
    subprocess.run([settings.ffmpeg_bin, '-y', '-f', 'lavfi', '-i', 'testsrc2=size=320x240:rate=10',
        '-t', '2', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', str(path)], check=True, capture_output=True, timeout=30)
    case, _ = new_case(client)
    upload_pair(client, case, path.read_bytes(), media='video', extension='mp4')
    job = client.post(f'/api/v1/cases/{case}/analyze').json()['job_id']
    status = client.get(f'/api/v1/jobs/{job}').json()
    assert status['status'] == 'completed', status
    report = client.get(f'/api/v1/cases/{case}/report').json()['report']
    assert report['headline_overlap_percentage'] == 100
    assert report['component_scores']['V4_transcript_similarity'] is None
    assert len(report['evidence']['timeline_matches']) == 4
    assert 'original_frame_path' not in report['evidence']['timeline_matches'][0]
    assert 'ffmpeg' in report['dependencies']
    assert client.get(f'/api/v1/cases/{case}/report.pdf').content.startswith(b'%PDF-')


def test_chunked_body_limit(client):
    body = (b' ' * 100_000 for _ in range(3))
    response = client.post('/api/v1/cases', content=body, headers={'Content-Type': 'application/json'})
    assert response.status_code == 413


def test_pdf_missing_glyphs_are_preserved_as_codepoints():
    from app.services.reports.pdf_renderer import portable_text
    assert portable_text('café') == 'café'
    assert portable_text('你好') == '[U+4F60][U+597D]'
